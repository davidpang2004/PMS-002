# -*- coding: utf-8 -*-
"""Generates Gemini_API_Key_设置指南.docx -- step-by-step guide for getting a
free (and later, paid) Gemini API key for the AI-assisted extraction feature.

Run with: python3 generate_gemini_guide.py
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.left_margin = section.right_margin = Inches(1.2)
section.top_margin = section.bottom_margin = Inches(1)

FONT = 'Microsoft YaHei'
PURPLE1 = RGBColor(0x6D, 0x28, 0xD9)
PURPLE2 = RGBColor(0x7C, 0x3A, 0xED)


def _cn(run):
    run.font.name = FONT
    try:
        run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    except Exception:
        pass


def h1(text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        _cn(r)
        r.font.color.rgb = PURPLE1


def h2(text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        _cn(r)
        r.font.color.rgb = PURPLE2


def para(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    _cn(r)
    return p


def numbered(text):
    p = doc.add_paragraph(style='List Number')
    r = p.add_run(text)
    _cn(r)


def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(text)
    _cn(r)


def tip(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(0.4)
    r1 = p.add_run('提示：')
    _cn(r1)
    r1.bold = True
    r1.font.color.rgb = RGBColor(0x05, 0x6F, 0x00)
    r2 = p.add_run(text)
    _cn(r2)
    r2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'E6F4EA')
    p._p.get_or_add_pPr().append(shd)


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = p.paragraph_format.right_indent = Inches(0.4)
    r1 = p.add_run('注意：')
    _cn(r1)
    r1.bold = True
    r1.font.color.rgb = RGBColor(0x92, 0x4E, 0x00)
    r2 = p.add_run(text)
    _cn(r2)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'FFF8E1')
    p._p.get_or_add_pPr().append(shd)


# ---------------------------------------------------------------------------

title = doc.add_heading('Gemini API Key 获取与设置指南', level=0)
for r in title.runs:
    _cn(r)
    r.font.color.rgb = PURPLE1

para('本指南说明如何获取 Google Gemini 的 API Key，用于 DMS 的"AI 辅助提取"功能'
     '（在"提取关键参数"面板中可选启用）。此功能默认关闭，不影响应用的离线使用。')

h1('一、获取免费 API Key')

numbered('打开浏览器，访问：aistudio.google.com/apikey')
numbered('使用 Google 账号登录（没有的话可以免费注册一个）')
numbered('点击"Create API key"（创建 API 密钥）')
numbered('选择创建到一个新项目，或使用已有的 Google Cloud 项目')
numbered('复制生成的密钥（以 "AIza" 开头的一串字符）')
numbered('回到 DMS，打开【项目菜单 → AI 辅助提取设置】，粘贴该密钥并点击"保存"')

tip('免费额度对偶尔使用完全够用，无需绑定信用卡。密钥保存后会加密存储在本机，'
    '不会显示在界面上，也不会被应用发送到 Google 以外的任何地方。')

note('免费额度有速率限制（每分钟/每天可调用的次数上限）。如果短时间内连续提取'
     '很多份文档，可能会遇到"暂时超出限额"的提示，稍后重试即可。')

h1('二、升级为付费额度（可选，将来需要时再做）')

para('升级不需要重新创建密钥——只需为同一个密钥所在的项目开通结算，'
     '密钥会自动获得更高的调用额度。')

numbered('再次打开 aistudio.google.com/apikey')
numbered('在项目列表中找到你的项目，"Billing Tier"（结算层级）会显示为 "Free Tier"')
numbered('点击该项目旁的"Set up billing"（设置结算）')
numbered('添加付款方式：')
bullet('新建结算账号：选择所在国家/地区，同意条款，填写联系方式和付款方式')
bullet('已有 Google Cloud 结算账号：直接选择使用')
numbered('选择付费方式：预付（Prepay，最低预存 $10）或后付（Postpay，如果系统提供此选项）')
numbered('完成设置后，DMS 中已保存的同一个密钥会自动获得付费层级的更高额度，无需在应用中做任何改动')

note('付费层级不仅仅是提高额度：根据 Google 的条款，免费层级的使用数据可能被用于'
     '改进其模型，而付费层级不会。如果提取的文档涉及隐私内容（如医疗、保险等），'
     '这也是提前考虑付费层级的一个原因。')

h1('三、常见问题')

h2('密钥可以随时更换或删除吗？')
para('可以。在【AI 辅助提取设置】对话框中输入新密钥并保存即可覆盖旧密钥；'
     '点击"移除 Key"可清除已保存的密钥，恢复为未配置状态。')

h2('不配置密钥会影响其他功能吗？')
para('不会。AI 辅助提取是完全可选的功能，未配置或未联网时，"提取关键参数"'
     '面板的默认关键字匹配方式照常离线工作，不受任何影响。')

out_path = HERE / "Gemini_API_Key_设置指南.docx"
doc.save(str(out_path))
print(f"Saved: {out_path}")
